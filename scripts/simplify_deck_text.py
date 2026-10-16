"""Rewrite the slide text of the final-review deck in simpler words.

Layout, colours and shapes are untouched: only the text of existing runs and
table cells is replaced. A paragraph that has two runs (bold lead + rest) is
given as a tuple so the bold lead stays bold.
"""

import sys
from pathlib import Path

from pptx import Presentation

DECK = Path(sys.argv[1])
p = Presentation(str(DECK))

# (slide, shape_id, paragraph_index) -> text  |  (lead, rest) for two-run paragraphs
TEXT = {
    # ---- 3 introduction
    (3, 21, 1): "•  Money is stuck in the warehouse",
    (3, 21, 2): "•  Storage costs keep adding up",
    (3, 21, 3): "•  Seasonal goods lose value",
    (3, 23, 1): "•  The customer goes elsewhere",
    (3, 23, 2): "•  The sale is lost for good",
    (3, 23, 3): "•  We also pay a penalty",
    (3, 25, 1): "•  Cheap but slow",
    (3, 25, 2): "•  Fast but costly",
    (3, 25, 3): "•  Looks fine but ships only 80%",
    (3, 26, 0): (
        "•  A chain of decisions. ",
        "Today’s order becomes next week’s stock, so every choice affects the next ones.",
    ),
    (3, 26, 1): (
        "•  Old rules fall short. ",
        "Textbook rules (like (s,S) and EOQ) look at each product alone and expect steady demand. In our data, daily demand swings about 216× more than a simple model expects.",
    ),
    (3, 26, 2): (
        "•  Our idea. ",
        "A reinforcement learning (RL) agent can learn when to order and from whom, and adapt to the season and to supplier problems, which a fixed formula cannot.",
    ),
    (3, 26, 3): (
        "•  Trust matters. ",
        "A manager will not follow a number they cannot understand. So a language model (LLM) explains each order in plain words, and every number it quotes is checked.",
    ),
    # ---- 4 problem definition (text boxes; tables below)
    (
        4,
        22,
        0,
    ): "No supplier is best: the cheap one is slow, the fast one costs more, and the middle one is unreliable.",
    (4, 24, 1): "different order combinations every day",
    (4, 26, 1): "how much more daily demand swings than a simple model expects",
    (4, 28, 1): "quantity and supplier are chosen together",
    # ---- 6 gaps
    (6, 21, 0): "Problems are too simple",
    (
        6,
        21,
        1,
    ): "Most studies use one product or a simple chain. Choosing quantity and supplier together, with one shared warehouse, is rarely studied.",
    (
        6,
        23,
        1,
    ): "We model 10 products, 3 suppliers and one shared warehouse, and choose both together.",
    (6, 25, 0): "No explanations",
    (6, 25, 1): "A learned agent gives a number, not a reason, so it stays in simulation.",
    (
        6,
        27,
        1,
    ): "An LLM explains each decision using facts from the simulator, and every number is checked.",
    (6, 29, 0): "Stress tests are made by hand",
    (6, 29, 1): "Crises are coded one by one, so only a few are ever tested.",
    (6, 31, 1): "An LLM writes crises that overlap; each is checked and kept within safe ranges.",
    (6, 33, 0): "Weak comparisons",
    (6, 33, 1): "Old rules are often left untuned, so a new method looks good for free.",
    (6, 35, 1): "We tuned all four classical methods first, on separate test runs.",
    # ---- 7 objectives
    (7, 21, 0): (
        "Objective 5 says compare, not win. ",
        "The result went against the AI, and we say so. What we can defend: RL reaches 90.8% of a strong, well-tuned baseline, and we know why it falls short.",
    ),
    # ---- 9 methodology
    (9, 22, 0): "Why masking",
    (9, 22, 1): "•  Masking blocks impossible choices first",
    (9, 22, 2): "•  Score: ₹24.9 lakh → ₹76.4 lakh (3×)",
    (9, 24, 0): "Why one joint choice",
    (9, 24, 1): "•  Quantity and supplier chosen together",
    (9, 24, 2): "•  Trains 2.8× faster, fits real life",
    (9, 26, 1): "•  Tried 8 settings, ran the best 2 longer",
    (9, 26, 2): "•  A slower learning rate worked best",
    # ---- 10 dataset
    (10, 21, 0): "Three kinds of demand: steady, moderate, bursty",
    (10, 22, 1): "•  1,067,371 sales records from a UK online gift shop",
    (
        10,
        22,
        2,
    ): "•  Cleaned: cancellations, returns and bad rows removed; one row per product per day; Saturdays dropped (the shop never opened)",
    (
        10,
        22,
        3,
    ): "•  10 products picked on purpose: 4 steady, 3 moderate, 3 bursty (very spiky sales)",
    (
        10,
        22,
        4,
    ): "•  The agent never sees the raw sales. We measure five things (average, ups and downs, weekday pattern, month pattern, zero-sale days) and the simulator recreates them",
    (
        10,
        23,
        0,
    ): "If all ten products behaved the same, one simple rule would work for all and an AI would have nothing to learn.",
    # ---- 11 baselines
    (11, 22, 0): (
        "•  Tuning mattered a lot. ",
        "Better settings lifted the best classical method from ₹69.7 lakh to ₹89.2 lakh. Without tuning, the AI would have looked better than it really is.",
    ),
    (11, 22, 1): (
        "•  Same choices for all. ",
        "Every classical method picks from the same 21 options per product as the AI.",
    ),
    (11, 22, 2): (
        "•  Same customers for all. ",
        "Every method faces the same customers and delays on each test run, so we compare them run by run.",
    ),
    # ---- 12 explainer
    (12, 21, 0): (
        "The rule: ",
        "the LLM is never asked why. The agent’s real reason is hidden inside a neural network, so any “because” from the LLM would just be a believable story.",
    ),
    # ---- 13 crises
    (
        13,
        21,
        0,
    ): "These events overlap: a supplier fails during the demand spike. A one-event test never checks that combination.",
    (13, 23, 0): "We do not trust what the LLM writes",
    (13, 23, 1): "•  The event type must be one we know",
    (13, 23, 2): "•  Suppliers and products must exist (wrong ones are dropped)",
    (13, 23, 3): "•  Numbers are kept within safe ranges",
    (13, 23, 4): "•  Events are moved to fit inside the 180 days",
    (13, 23, 5): "•  6 deliberately bad inputs: none got through as written",
    (13, 25, 0): "Training on crises",
    (13, 25, 1): "•  16 sets of crises (from 19 LLM scenarios): 10 for training, 6 for testing",
    (
        13,
        25,
        2,
    ): "•  The agent trains on the 10 and is tested on the 6 it never saw. Everything else matches the plain agent, so crises are the only difference",
    (
        13,
        25,
        3,
    ): "•  The split matters: testing on crises it trained on would only prove it memorised them",
    # ---- 14 dashboard 1
    (14, 21, 0): "Dashboard: forecast + safety stock policy, test run 500",
    (14, 23, 0): "Choose the policy and test run",
    (14, 23, 1): "Sidebar: any policy, any test run, with or without a crisis.",
    (14, 25, 1): "Profit ₹87.8 lakh, share of demand served, order fees, lost sales, wasted goods.",
    (14, 27, 1): "Two simple charts, each with its own scale.",
    # ---- 15 dashboard 2
    (15, 21, 0): "RL challenger: orders spread across all three suppliers",
    (15, 23, 0): "Engine with a crisis added: shaded band = disruption",
    # ---- 16 dashboard 3
    (
        16,
        21,
        0,
    ): "“Before”: the same fixed order every day.  “After”: our engine, with the RL challenger beside it.",
    (16, 25, 1): "Share of demand served   +5.6 points",
    (16, 27, 1): "Lost-sales cost   -79%",
    (
        16,
        28,
        0,
    ): "Most of this gain comes from smart forecasting with safety stock (the engine), not from RL.",
    # ---- 17 dashboard 4
    (17, 21, 0): "Day 45, product 0: explanation written by the LLM; the badge shows who wrote it",
    (17, 22, 1): "•  Choose a day and a product",
    (
        17,
        22,
        4,
    ): "•  The black badge shows the writer: a model name, or “template” if our check rejected the LLM",
    (17, 22, 5): "•  Open “the facts the model was given” to check every number",
    (17, 23, 0): (
        "Check it yourself: ",
        "2,365 units, 14.0 days of cover, demand of 169 units and a 7-day average of 78 all appear in the facts panel. If the model had made up a number, the template would have replaced this text.",
    ),
    # ---- 18 dashboard 5
    (18, 21, 0): "Comparison tab: 30 unseen test runs",
    (18, 23, 0): "Robustness tab: normal days vs 6 unseen crises",
    (
        18,
        24,
        0,
    ): "•  Only the rows needed for our claims are shown: the 4 classical methods and 3 RL agents (tuned, untuned, no masking); then the engine, the tuned agent, and the crisis-trained agent beside its control.",
    (
        18,
        24,
        1,
    ): "•  The cards give the key numbers; the table under each chart has the same data for checking.",
    # ---- 19 results 1
    (19, 25, 1): "gap per run (t = −6.45: a real gap)",
    (
        19,
        26,
        0,
    ): "The best classical method beats the tuned agent, and the gap is real, not luck (95% range: 5.5–10.6 lakh). The agent does beat two of the four classical methods: EOQ and (s,S).",
    (19, 26, 1): "Average profit per 180-day run over 30 unseen test runs. Rupees at 1 GBP = 105.",
    # ---- 20 results 2
    (20, 21, 0): (
        "•  Masking: about 3×. ",
        "Without it, the agent wastes most of its training on choices that can never work.",
    ),
    (20, 21, 1): (
        "•  Tuning: closed 38% of what was left. ",
        "The gap shrank from 13.0 to 8.1 lakh.",
    ),
    (20, 21, 2): "Practice scores use runs 900–905; test scores use unseen runs 500–529.",
    (20, 23, 0): "Learning curves: no masking, masking, tuned",
    # ---- 21 results 3
    (21, 22, 1): "•  In a crisis it earns ₹38.9 lakh, same as the control’s ₹38.8 lakh",
    (21, 22, 2): "•  and it lost ₹7.7 lakh of normal profit",
    (21, 22, 3): "•  t = +1.69: not a real difference",
    (21, 23, 0): (
        "•  A smaller % drop can mislead ",
        "when the starting profit is lower: the crisis-trained agent drops less because it had less to lose. We show the drop next to the actual profit so we are not fooled.",
    ),
    (21, 23, 1): (
        "•  The useful finding: ",
        "simply tuning the settings made the agent more crisis-proof than crisis training did. The tuned agent never trained on a crisis, yet has the highest crisis profit and the smallest drop of the agents.",
    ),
    (
        21,
        23,
        2,
    ): "Limit: the crisis-trained agent had a harder job with the same training time, so it may simply be under-trained.",
    # ---- 22 results 4
    (22, 21, 1): "raw LLM accuracy: correct and complete, before our check",
    (22, 23, 1): "delivered accuracy: what the manager actually saw",
    (22, 25, 1): "outputs had an invented number; all 5 caught",
    (22, 27, 1): "•  The five bad outputs sounded confident and fluent. A buyer would not notice.",
    (22, 27, 2): "•  Without the check, 1 explanation in 6 would contain a made-up number.",
    (
        22,
        27,
        3,
    ): "•  An earlier version only scored outputs that had already passed the check, so it showed 100% for a model that made up numbers 17% of the time. We now show both numbers.",
    # ---- 23 conclusion
    (
        23,
        21,
        1,
    ): "•  A complete system: simulator, RL agent, 4 tuned classical methods, LLM explainer, crisis generator, dashboard",
    (
        23,
        21,
        2,
    ): "•  The tuned agent earns 90.8% of the best classical method and beats two of the four; the gap is real (t = -6.45) and comes from splitting orders across suppliers",
    (23, 21, 4): "•  The explainer is accurate only because it is checked: 83% before, 100% after",
    (
        23,
        21,
        5,
    ): "•  Safe way to use it: the classical engine decides, RL is scored alongside, nothing is executed",
    (
        23,
        23,
        1,
    ): "•  Give the agent a memory of its own orders (days since last order, last supplier, fees so far) and test again on the same 30 runs: the best chance to close the gap",
    (23, 23, 2): "•  Test all methods when the market changes (wilder demand, longer delays)",
    (
        23,
        23,
        3,
    ): "•  Train longer on crises, to see whether the method failed or was just under-trained",
    (23, 23, 4): "•  Use recent real data and real supplier terms if a partner can provide them",
    (
        23,
        23,
        5,
    ): "•  Before real use: handle new products, set order limits, keep an audit trail, and flag when a human should decide",
}

# slide titles
TITLES = {
    11: "Classical Methods and a Fair Test",
    13: "LLM Crisis Generator and Training on Crises",
    21: "Results (3/4): Coping with Unseen Crises",
}

# (slide, shape_id) -> list of rows; a row is a list of cell texts, None = leave as is
TABLES = {
    (4, 20): [
        ["What", "Details"],
        ["Products", "10 real products from a UK shop’s sales data"],
        ["Suppliers", "3, made up by us (no public data on suppliers)"],
        ["Time", "180 days per run; the first 14 days are a warm-up and not counted"],
        ["Warehouse", "one shared store, 20,000 units; extra goods are turned away"],
        ["Decision", "every day, for each product: 7 quantities × 3 suppliers = 21 choices"],
        [
            "Goal",
            "maximise profit = sales − purchases − storage − fees − lost sales − wasted goods",
        ],
    ],
    (4, 21): [["Supplier", "Price", "Delivery", "Delivers", "Fee"], None, None, None],
    (5, 20): [
        None,
        [
            None,
            "Proximal Policy Optimization (2017)",
            "A stable way to train an RL agent",
            "Reliable; a common default",
            "General purpose; not built for supply chains",
        ],
        [
            None,
            "A Deep Q-Network for the Beer Game (2022)",
            "Deep RL on a simple supply chain",
            "Matches or beats standard ordering rules when teams cooperate",
            "One product, no supplier choice",
        ],
        [
            None,
            "Can Deep RL Improve Inventory Management? (2022)",
            "Deep RL on lost sales and two suppliers",
            "As good as well-tuned rules, not clearly better",
            "Needs a lot of tuning; no explanations",
        ],
        [
            None,
            "Deep RL for Inventory Control: A Roadmap (2022)",
            "Survey of the field",
            "Says fair comparisons and explanations are still open problems",
            "Ideas only; no experiments",
        ],
        [
            None,
            "A Closer Look at Invalid Action Masking (2022)",
            "Blocking impossible actions during training",
            "Works well even with very many actions",
            "Not tried on inventory",
        ],
    ],
    (7, 20): [
        None,
        [
            None,
            "Build a supply-chain simulation from real data",
            None,
            "Built from 1.07M real sales; sensible ordering earns profit, silly ordering loses",
        ],
        [
            None,
            "Train an RL agent to choose quantity and supplier",
            None,
            "Tuned agent, 2M training steps; earns 90.8% of the best classical method",
        ],
        [
            None,
            "LLM explains each decision in plain business language",
            None,
            "100% of explanations shown were correct; 83% before our check, which caught every invented number",
        ],
        [
            None,
            "LLM creates demand-spike and supplier-delay scenarios",
            None,
            "6 unseen crises used for testing; crisis training also tried",
        ],
        [
            None,
            "Compare the RL agent with traditional policies",
            None,
            "30 unseen test runs: the classical method still wins",
        ],
    ],
    (11, 20): [
        None,
        [None, "order when stock covers under 11 days; refill up to 16 days"],
        [None, "order at 8 days of stock left, with a safety margin of 1.6"],
        [None, "plan 3 days ahead, use the cheapest supplier"],
        [
            None,
            "forecast from yesterday, extra safety stock (about 90% service), check every 4 days",
        ],
    ],
    (11, 21): [
        ["Seeds (test runs)", "Used for"],
        [None, "tuning the classical methods"],
        [None, "choosing the best RL version"],
        [None, "final results only; never used to make any choice"],
    ],
    (13, 20): [
        None,
        [None, "demand triples on four products", None],
        [None, "supplier 2 goes offline", None],
        [None, None, None],
        [None, "costs rise 1.6×", None],
    ],
    (19, 20): [
        ["Method", "Profit", "Served", "Fees", "Lost sales"],
        None,
        None,
        None,
        None,
        None,
        None,
        ["Fixed daily order (no system)", None, None, None, None],
        None,
    ],
    (20, 20): [
        ["Step", "Practice score", "Test score", "Gap to classical"],
        ["Basic RL (no masking)", None, None, None],
        ["+ masking (block bad choices)", None, None, None],
        None,
    ],
    (21, 20): [
        ["Method", "Normal", "Crisis", "Drop"],
    ],
    (22, 26): [
        None,
        ["LLM answered (after retries)", None],
        ["Only real numbers used (before the check)", None],
    ],
}


def set_single(para, text):
    runs = para.runs
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def set_runs(para, parts):
    runs = para.runs
    if len(runs) >= 2:
        runs[0].text, runs[1].text = parts
        for r in runs[2:]:
            r.text = ""
    else:
        runs[0].text = parts[0] + parts[1]


changed = missing = 0
for idx, slide in enumerate(p.slides, 1):
    by_id = {sh.shape_id: sh for sh in slide.shapes}
    for (sn, sid, pi), new in TEXT.items():
        if sn != idx:
            continue
        sh = by_id.get(sid)
        if sh is None or not sh.has_text_frame or pi >= len(sh.text_frame.paragraphs):
            print("MISSING", sn, sid, pi)
            missing += 1
            continue
        para = sh.text_frame.paragraphs[pi]
        if isinstance(new, tuple):
            set_runs(para, new)
        else:
            set_single(para, new)
        changed += 1
    for (sn, sid), rows in TABLES.items():
        if sn != idx:
            continue
        sh = by_id.get(sid)
        if sh is None or not sh.has_table:
            print("MISSING table", sn, sid)
            missing += 1
            continue
        for ri, row in enumerate(rows):
            if row is None:
                continue
            for ci, val in enumerate(row):
                if val is None:
                    continue
                cell = sh.table.cell(ri, ci)
                set_single(cell.text_frame.paragraphs[0], val)
                changed += 1
    if idx in TITLES:
        t = slide.shapes.title
        set_single(t.text_frame.paragraphs[0], TITLES[idx])
        changed += 1

p.save(str(DECK))
print(f"edited {changed} text items; {missing} not found")
